from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
SOURCE_MD = ROOT / "InvoiceCraft_SI_Report_Content.md"
TEMPLATE_DIR = Path(r"C:\Users\dhruv\Downloads\SI_Report_Format-20260804T100103Z-1-001\SI_Report_Format")
OUTPUT_DIR = ROOT / "InvoiceCraft_Report_Docs"

TITLE = "InvoiceCraft: A Client-Side Web Application for Professional Invoice Generation and PDF Export"


def md_section(text: str, start: str, end: str | None = None) -> str:
    lines = text.splitlines()
    try:
        start_i = lines.index(start)
    except ValueError as exc:
        raise RuntimeError(f"Missing markdown heading: {start}") from exc
    end_i = len(lines)
    if end is not None:
        try:
            end_i = lines.index(end)
        except ValueError as exc:
            raise RuntimeError(f"Missing markdown end heading: {end}") from exc
    return "\n".join(lines[start_i + 1 : end_i]).strip()


def clean_markup(value: str) -> str:
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"__(.*?)__", r"\1", value)
    value = value.replace("`", "")
    value = value.replace("***", "")
    value = value.replace("—", "-").replace("–", "-")
    return value.strip()


def markdown_blocks(raw: str):
    """Return (kind, text) blocks suitable for Word paragraphs."""
    blocks = []
    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    in_code = False

    def flush_paragraph():
        nonlocal paragraph_lines
        if paragraph_lines:
            text = clean_markup(" ".join(x.strip() for x in paragraph_lines))
            if text:
                blocks.append(("paragraph", text))
            paragraph_lines = []

    def flush_code():
        nonlocal code_lines
        if code_lines:
            text = "\n".join(code_lines).strip()
            if text:
                blocks.append(("code", text))
            code_lines = []

    for line in raw.splitlines():
        stripped = line.strip()
        if stripped == "```":
            flush_paragraph()
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line.rstrip())
            continue
        if not stripped or stripped == "---":
            flush_paragraph()
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            blocks.append(("heading3", clean_markup(stripped[4:])))
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            blocks.append(("heading2", clean_markup(stripped[3:])))
            continue
        if stripped.startswith("# "):
            flush_paragraph()
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            blocks.append(("bullet", clean_markup(stripped[2:])))
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            flush_paragraph()
            blocks.append(("number", clean_markup(numbered.group(1))))
            continue
        if stripped.startswith("|"):
            # Keep report tables readable without depending on a particular table style.
            flush_paragraph()
            if set(stripped.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                continue
            cells = [clean_markup(c.strip()) for c in stripped.strip("|").split("|")]
            blocks.append(("table_row", "    ".join(cells)))
            continue
        paragraph_lines.append(stripped)

    flush_paragraph()
    flush_code()
    return blocks


def all_paragraphs(doc: Document):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    return paragraphs


def set_paragraph_text(paragraph, text: str):
    paragraph.text = text


def replace_text_everywhere(doc: Document, replacements: dict[str, str]):
    for paragraph in all_paragraphs(doc):
        current = paragraph.text
        updated = current
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != current:
            set_paragraph_text(paragraph, updated)


def find_paragraph(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise RuntimeError(f"Could not find paragraph: {text}")


def has_page_break(paragraph) -> bool:
    xml = paragraph._p.xml
    return 'w:type="page"' in xml or "lastRenderedPageBreak" in xml


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def remove_blank_paragraphs_between(doc: Document, start_text: str, end_text: str):
    start = find_paragraph(doc, start_text)
    end = find_paragraph(doc, end_text)
    paragraphs = list(doc.paragraphs)
    start_i = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is start._p)
    end_i = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is end._p)
    for paragraph in paragraphs[start_i + 1 : end_i]:
        if not paragraph.text.strip() and not has_page_break(paragraph):
            remove_paragraph(paragraph)


def add_blocks_before(anchor, raw: str):
    for kind, text in markdown_blocks(raw):
        requested_style = {
            "heading2": "Heading 2",
            "heading3": "Heading 3",
            "bullet": "List Bullet",
            "number": "List Number",
            "code": "No Spacing",
        }.get(kind, "Normal")
        style_names = {style.name for style in anchor._parent.part.styles}
        style = requested_style if requested_style in style_names else "Normal"
        paragraph = anchor.insert_paragraph_before(text, style=style)
        if kind.startswith("heading"):
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(13 if kind == "heading2" else 11)
        if kind == "code":
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
                run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")


def remove_blank_paragraphs_after(doc: Document, start_text: str):
    start = find_paragraph(doc, start_text)
    paragraphs = list(doc.paragraphs)
    start_i = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is start._p)
    for paragraph in paragraphs[start_i + 1 :]:
        if not paragraph.text.strip() and not has_page_break(paragraph):
            remove_paragraph(paragraph)


def add_blocks_at_end(doc: Document, raw: str):
    for kind, text in markdown_blocks(raw):
        requested_style = {
            "heading2": "Heading 2",
            "heading3": "Heading 3",
            "bullet": "List Bullet",
            "number": "List Number",
            "code": "No Spacing",
        }.get(kind, "Normal")
        style_names = {style.name for style in doc.styles}
        style = requested_style if requested_style in style_names else "Normal"
        paragraph = doc.add_paragraph(text, style=style)
        if kind.startswith("heading"):
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(13 if kind == "heading2" else 11)
        if kind == "code":
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
                run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")


def add_heading_fallback(doc: Document, text: str, level: int = 2):
    requested_style = f"Heading {level}"
    style_names = {style.name for style in doc.styles}
    style = requested_style if requested_style in style_names else "Normal"
    paragraph = doc.add_paragraph(text, style=style)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(13 if level == 2 else 11)
    return paragraph


def replace_front_matter(source_text: str):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Title page
    title_src = TEMPLATE_DIR / "0 - Title page.docx"
    title_out = OUTPUT_DIR / "00 - Title page - InvoiceCraft.docx"
    shutil.copy2(title_src, title_out)
    doc = Document(title_out)
    replace_text_everywhere(doc, {
        "Title of the Project": TITLE,
        "Name of Student.": "[STUDENT NAME]",
        "(Enrollment no.)": "([ENROLLMENT NUMBER])",
        "Faculty Name": "[GUIDE NAME]",
    })
    doc.save(title_out)

    # Certificate
    src = TEMPLATE_DIR / "1-CERTIFICATE.docx"
    out = OUTPUT_DIR / "01 - Certificate - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replace_text_everywhere(doc, {
        "PROJECT TITLE": TITLE,
        "STUDENT NAME": "[STUDENT NAME]",
        "Signature and Name of Guide": "[GUIDE NAME]",
    })
    # Keep the template's guide block readable after replacing its name slot.
    guide_cell = doc.tables[0].cell(0, 0)
    if len(guide_cell.paragraphs) >= 2:
        guide_cell.paragraphs[1].text = ""
    doc.save(out)

    # Declaration
    src = TEMPLATE_DIR / "2-DECLARATION- Online & Ofline.docx"
    out = OUTPUT_DIR / "02 - Declaration - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replace_text_everywhere(doc, {
        "PROJECT TITLE": TITLE,
        "Guide Name": "[GUIDE NAME]",
    })
    # The student name is a separate table slot; avoid replacing NAME inside
    # the already-filled [GUIDE NAME] placeholder.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip() == "NAME":
                        paragraph.text = "[STUDENT NAME]"
    doc.save(out)

    # Acknowledgement
    src = TEMPLATE_DIR / "3-ACKNOWLEDGEMENT.docx"
    out = OUTPUT_DIR / "03 - Acknowledgement - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    for paragraph in doc.paragraphs:
        if "Project Guide" in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace("NAME", "[GUIDE NAME]"))
        elif "H.O.D." in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace("NAME", "Dr. Gayatri Pandi"))
    doc.save(out)

    # Abstract
    src = TEMPLATE_DIR / "4 - Abstract.docx"
    out = OUTPUT_DIR / "04 - Abstract - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replace_text_everywhere(doc, {
        "Title of the Project": TITLE,
        "Enrollment No.:": "Enrollment No.: [ENROLLMENT NUMBER]",
        "Student Name:": "Student Name: [STUDENT NAME]",
        "*** For Example***": "",
    })
    abstract_paragraph = (
        "InvoiceCraft is a responsive, browser-based invoice generation application designed for freelancers, "
        "students, consultants, and small businesses that need a quick method of preparing professional invoices. "
        "Traditional invoice preparation using word processors or spreadsheets often requires repeated manual formatting "
        "and calculation, which can lead to arithmetic mistakes, inconsistent layouts, and loss of draft data. InvoiceCraft "
        "addresses these problems through a guided editor and a live invoice preview. The application is implemented using "
        "React and TypeScript with Vite as the build tool. Users can enter company and client information, add multiple line "
        "items, specify HSN/SAC codes, select a currency, apply tax and discounts, record the amount paid, and view the balance "
        "due. It also supports payment information, optional UPI QR-code generation, signatures, notes, terms and conditions, "
        "brand colors, text colors, and multiple invoice templates. Invoice data is stored in the browser's LocalStorage so "
        "drafts can be recovered without requiring an account or server-side database. The completed invoice is converted into "
        "a high-quality A4 PDF using html2canvas and jsPDF. The project demonstrates how a modern client-side application can "
        "combine reusable components, state management, responsive design, automatic calculations, and document export to solve "
        "a practical business problem."
    )
    # The template's example paragraph is the final non-empty paragraph in the body.
    body_paragraphs = doc.paragraphs
    for paragraph in reversed(body_paragraphs):
        if paragraph.text.strip() == "":
            continue
        if paragraph.text.startswith("Python is a widely") or paragraph.text == "":
            set_paragraph_text(paragraph, abstract_paragraph)
            break
    doc.add_paragraph("Keywords: invoice generator, React, TypeScript, client-side application, LocalStorage, PDF export, GST, UPI QR code, responsive web design.")
    doc.save(out)

    # Table of contents
    src = TEMPLATE_DIR / "5-TABLE OF CONTENT.docx"
    out = OUTPUT_DIR / "05 - Table of Contents - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replacements = {
        "CHAPTER 1 INTRODUCTION": "CHAPTER 1 INTRODUCTION",
        "2.2 Introduction to Python Programming": "2.2 Introduction to React and TypeScript",
        "2.2 Key Frameworks and Libraries": "2.3 Key Frameworks and Libraries",
    }
    replace_text_everywhere(doc, replacements)
    doc.save(out)


def replace_chapter_1(source_text: str):
    src = TEMPLATE_DIR / "CHAPTER 1 INTRODUCTION (1-4).docx"
    out = OUTPUT_DIR / "06 - Chapter 1 Introduction - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)

    headings = {
        "Project Introduction": "1.1 Project Introduction",
        "Problem Statement and Objective": "1.2 Problem Statement and Objective",
        "Scope of the Project": "1.3 Scope of the Project",
        "1.4 Work Plan and Execution": "1.4 Work Plan and Execution",
    }
    replace_text_everywhere(doc, headings)
    s11 = md_section(source_text, "## 1.1 Project Introduction", "## 1.2 Problem Statement and Objective")
    s12 = md_section(source_text, "## 1.2 Problem Statement and Objective", "## 1.3 Scope of the Project")
    s13 = md_section(source_text, "## 1.3 Scope of the Project", "## 1.4 Work Plan and Execution")
    s14 = md_section(source_text, "## 1.4 Work Plan and Execution", "---")

    remove_blank_paragraphs_between(doc, "1.1 Project Introduction", "1.2 Problem Statement and Objective")
    add_blocks_before(find_paragraph(doc, "1.2 Problem Statement and Objective"), s11)
    remove_blank_paragraphs_between(doc, "1.2 Problem Statement and Objective", "1.3 Scope of the Project")
    add_blocks_before(find_paragraph(doc, "1.3 Scope of the Project"), s12)
    remove_blank_paragraphs_between(doc, "1.3 Scope of the Project", "1.4 Work Plan and Execution")
    # Preserve the template's page break before section 1.4.
    page_break_anchor = None
    for paragraph in doc.paragraphs:
        if has_page_break(paragraph) and not paragraph.text.strip():
            page_break_anchor = paragraph
            break
    if page_break_anchor is None:
        page_break_anchor = find_paragraph(doc, "1.4 Work Plan and Execution")
    add_blocks_before(page_break_anchor, s13)
    remove_blank_paragraphs_after(doc, "1.4 Work Plan and Execution")
    add_blocks_at_end(doc, s14)
    doc.save(out)


def replace_chapter_2(source_text: str):
    src = TEMPLATE_DIR / "CHAPTER 2 TOOLS AND TECHNOLOGIES USED (5-7).docx"
    out = OUTPUT_DIR / "07 - Chapter 2 Tools and Technologies - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replace_text_everywhere(doc, {
        "2.2 Introduction to Python Programming": "2.2 Introduction to React and TypeScript",
        "2.3 Key Frameworks and Libraries": "2.3 Key Frameworks and Libraries",
    })
    s21 = md_section(source_text, "## 2.1 Hardware and Software Requirements", "## 2.2 Introduction to React and TypeScript")
    s22 = md_section(source_text, "## 2.2 Introduction to React and TypeScript", "## 2.3 Key Frameworks and Libraries")
    s23 = md_section(source_text, "## 2.3 Key Frameworks and Libraries", "# CHAPTER 3 SYSTEM DESIGN AND IMPLEMENTATION")
    remove_blank_paragraphs_after(doc, "2.1 Hardware and Software Requirements")
    add_blocks_at_end(doc, s21)
    add_heading_fallback(doc, "2.2 Introduction to React and TypeScript", 2)
    add_blocks_at_end(doc, s22)
    add_heading_fallback(doc, "2.3 Key Frameworks and Libraries", 2)
    add_blocks_at_end(doc, s23)
    doc.save(out)


def replace_chapter_3(source_text: str):
    src = TEMPLATE_DIR / "CHAPTER 3 SYSTEM DESIGN AND IMPLEMENTATION (10-23).docx"
    out = OUTPUT_DIR / "08 - Chapter 3 System Design and Implementation - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    s31 = md_section(source_text, "## 3.1 System Architecture / Flow Diagram", "## 3.2 Modules Description")
    s32 = md_section(source_text, "## 3.2 Modules Description", "## 3.3 Implementation Details")
    s33 = md_section(source_text, "## 3.3 Implementation Details", "# CHAPTER 4 CONCLUSION")
    remove_blank_paragraphs_after(doc, "3.1 System Architecture / Flow Diagram")
    add_blocks_at_end(doc, s31)
    doc.add_page_break()
    add_heading_fallback(doc, "3.2 Modules Description", 2)
    add_blocks_at_end(doc, s32)
    doc.add_page_break()
    add_heading_fallback(doc, "3.3 Implementation Details", 2)
    add_blocks_at_end(doc, s33)
    doc.save(out)


def replace_chapter_4(source_text: str):
    src = TEMPLATE_DIR / "CHAPTER 4 Conclusion (24-25).docx"
    out = OUTPUT_DIR / "09 - Chapter 4 Conclusion - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replace_text_everywhere(doc, {"4.1 Conclusion:": "4.1 Conclusion"})
    s41 = md_section(source_text, "## 4.1 Conclusion", "# CHAPTER 5 FUTURE ENHANCEMENT")
    remove_blank_paragraphs_after(doc, "4.1 Conclusion")
    add_blocks_at_end(doc, s41)
    doc.save(out)


def replace_chapter_5(source_text: str):
    src = TEMPLATE_DIR / "CHAPTER 5  Future Enhancement (26-27).docx"
    out = OUTPUT_DIR / "10 - Chapter 5 Future Enhancement - InvoiceCraft.docx"
    shutil.copy2(src, out)
    doc = Document(out)
    replace_text_everywhere(doc, {"5.1 Future Enhancement:": "5.1 Future Enhancement"})
    s51 = md_section(source_text, "## 5.1 Future Enhancement", "# Final field checklist")
    remove_blank_paragraphs_after(doc, "5.1 Future Enhancement")
    add_blocks_at_end(doc, s51)
    doc.save(out)


def main():
    source_text = SOURCE_MD.read_text(encoding="utf-8")
    if not TEMPLATE_DIR.exists():
        raise RuntimeError(f"Template directory not found: {TEMPLATE_DIR}")
    replace_front_matter(source_text)
    replace_chapter_1(source_text)
    replace_chapter_2(source_text)
    replace_chapter_3(source_text)
    replace_chapter_4(source_text)
    replace_chapter_5(source_text)
    print(f"Created {len(list(OUTPUT_DIR.glob('*.docx')))} DOCX files in {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
